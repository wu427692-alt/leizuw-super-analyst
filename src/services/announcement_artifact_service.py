# -*- coding: utf-8 -*-
"""Excel/PDF/TXT artifacts backed by persisted CNInfo monitoring events."""

from __future__ import annotations

from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import threading
import time
from typing import Any, Dict, Iterable, List, Tuple
from urllib.parse import urlparse
import zipfile

from openpyxl import Workbook
import requests


class AnnouncementArtifactError(RuntimeError):
    """Safe artifact generation or upstream file error."""


class _PdfExtractionFailure(AnnouncementArtifactError):
    """Internal typed failure used to decide whether a bounded fallback is safe."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def _extract_pdf_text_worker(pdf_path: Path, engine: str, max_text_chars: int) -> Dict[str, Any]:
    """Extract every page for the isolated worker process.

    This function must never return a partial document as a successful result.
    The parent process validates the page counters again before caching it.
    """

    started = time.monotonic()
    pages: List[str] = []
    page_count = 0
    extracted_chars = 0
    engine_version = "unknown"
    try:
        if engine == "pdfium":
            try:
                import pypdfium2 as pdfium
            except ImportError as exc:
                raise _PdfExtractionFailure("dependency_unavailable", "缺少 pypdfium2") from exc
            engine_version = (
                f"pypdfium2={getattr(pdfium, 'V_PYPDFIUM2', 'unknown')};"
                f"pdfium={getattr(pdfium, 'V_PDFIUM', 'unknown')}"
            )
            document = pdfium.PdfDocument(str(pdf_path))
            try:
                page_count = len(document)
                for page_index in range(page_count):
                    page = document[page_index]
                    try:
                        text_page = page.get_textpage()
                        try:
                            page_text = text_page.get_text_bounded(errors="ignore")
                        finally:
                            text_page.close()
                    finally:
                        page.close()
                    normalized_text = (page_text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
                    pages.append(normalized_text)
                    extracted_chars += len(normalized_text)
                    if extracted_chars > max_text_chars:
                        raise _PdfExtractionFailure("text_limit", "PDF 提取文本超过安全上限")
            finally:
                document.close()
            extraction_method = "pdfium_text"
        elif engine == "pypdf":
            try:
                import pypdf
                from pypdf import PdfReader
            except ImportError as exc:
                raise _PdfExtractionFailure("dependency_unavailable", "缺少 pypdf") from exc
            engine_version = f"pypdf={getattr(pypdf, '__version__', 'unknown')}"
            reader = PdfReader(str(pdf_path))
            page_count = len(reader.pages)
            for page in reader.pages:
                normalized_text = (page.extract_text() or "").replace("\r\n", "\n").replace("\r", "\n").strip()
                pages.append(normalized_text)
                extracted_chars += len(normalized_text)
                if extracted_chars > max_text_chars:
                    raise _PdfExtractionFailure("text_limit", "PDF 提取文本超过安全上限")
            extraction_method = "pypdf_text_fallback"
        else:
            raise _PdfExtractionFailure("unsupported_engine", "未知 PDF 提取引擎")
    except _PdfExtractionFailure:
        raise
    except Exception as exc:
        raise _PdfExtractionFailure("parse_failed", f"{engine} 解析失败：{type(exc).__name__}") from exc

    pages_extracted = len(pages)
    if page_count <= 0 or pages_extracted != page_count:
        raise _PdfExtractionFailure("incomplete", "PDF 页数校验未通过")
    text = "\n\n".join(value for value in pages if value)
    return {
        "text": text,
        "page_count": page_count,
        "pages_extracted": pages_extracted,
        "extraction_complete": True,
        "extraction_status": "complete" if text else "complete_no_selectable_text",
        "extraction_method": extraction_method,
        "extraction_engine_version": engine_version,
        "duration_ms": max(0, round((time.monotonic() - started) * 1000)),
    }


def _write_pdf_worker_result(output_path: Path, payload: Dict[str, Any]) -> None:
    output_path.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


def _run_pdf_text_worker_cli(argv: List[str]) -> int:
    """Private subprocess entrypoint; stdout/stderr never contains document text."""

    if len(argv) != 5 or argv[1] != "--pdf-text-worker":
        return 64
    engine, pdf_value, output_value = argv[2], argv[3], argv[4]
    output_path = Path(output_value)
    try:
        max_text_chars = int(os.environ.get("DSA_PDF_WORKER_MAX_TEXT_CHARS", "12000000"))
        payload = {"status": "success", **_extract_pdf_text_worker(Path(pdf_value), engine, max_text_chars)}
        _write_pdf_worker_result(output_path, payload)
        return 0
    except _PdfExtractionFailure as exc:
        _write_pdf_worker_result(output_path, {"status": "failed", "error_code": exc.code})
        return 2
    except Exception:
        _write_pdf_worker_result(output_path, {"status": "failed", "error_code": "worker_crashed"})
        return 3


class AnnouncementArtifactService:
    _TEXT_CACHE_VERSION = 2
    _LEGACY_TEXT_CACHE_VERSION = 1
    _PDFIUM_TIMEOUT_SECONDS = 45
    _PYPDF_FALLBACK_TIMEOUT_SECONDS = 75
    _MAX_EXTRACTED_TEXT_CHARS = 12_000_000
    _FALLBACK_SAFE_FAILURES = frozenset({"dependency_unavailable", "parse_failed", "worker_crashed", "incomplete"})
    _SUPPORTED_EXTRACTION_METHODS = frozenset({"pdfium_text", "pypdf_text_fallback", "pypdf_text"})

    def __init__(self, session: requests.Session | None = None):
        self.session = session or requests.Session()
        self.root = Path(os.getenv("ANNOUNCEMENT_FILE_DIR", "./data/announcements")).expanduser().resolve()
        self.max_pdf_bytes = max(1, min(int(os.getenv("ANNOUNCEMENT_PDF_MAX_MB", "40")), 200)) * 1024 * 1024
        self._text_cache_lock = threading.RLock()
        self.root.mkdir(parents=True, exist_ok=True)

    def excel_bytes(self, events: Iterable[Dict[str, Any]]) -> bytes:
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("上市公司公告")
        sheet.append(["公告ID", "证券代码", "公司名称", "公告标题", "公告日期", "公告分类", "PDF原文", "文件大小KB"])
        for event in events:
            metrics = event.get("metrics") or {}
            sheet.append([
                event.get("external_id"), (event.get("symbols") or [metrics.get("code") or ""])[0],
                (event.get("actors") or [metrics.get("name") or ""])[0], event.get("title"),
                str(event.get("event_at") or "")[:10], " / ".join(event.get("tags") or []),
                event.get("url"), metrics.get("size_kb"),
            ])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def package(self, events: List[Dict[str, Any]], *, include_text: bool = True) -> Dict[str, Any]:
        if not events:
            raise AnnouncementArtifactError("没有可打包的已入库公告")
        if len(events) > 20:
            raise AnnouncementArtifactError("单次最多打包 20 份公告，请缩小日期或股票范围")
        handle = tempfile.NamedTemporaryFile(prefix="cninfo_announcements_", suffix=".zip", delete=False)
        archive_path = Path(handle.name)
        handle.close()
        results: List[Dict[str, Any]] = []
        try:
            with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("上市公司公告索引.xlsx", self.excel_bytes(events))
                for event in events:
                    result = self._archive_event(archive, event, include_text=include_text)
                    results.append(result)
                archive.writestr("manifest.json", json.dumps(results, ensure_ascii=False, indent=2))
        except Exception:
            archive_path.unlink(missing_ok=True)
            raise
        return {
            "path": archive_path, "filename": "上市公司公告_PDF_TXT.zip",
            "requested": len(events), "downloaded": sum(1 for item in results if item["status"] == "success"),
            "failed": sum(1 for item in results if item["status"] == "failed"), "items": results,
        }

    def _archive_event(self, archive: zipfile.ZipFile, event: Dict[str, Any], *, include_text: bool) -> Dict[str, Any]:
        metrics = event.get("metrics") or {}
        announcement_id = str(event.get("external_id") or metrics.get("announcement_id") or "announcement")
        symbol = str((event.get("symbols") or [metrics.get("code") or "unknown"])[0]).split(".", 1)[0]
        title = self._safe_name(str(event.get("title") or announcement_id), 100)
        prefix = f"{symbol}_{announcement_id}_{title}"
        try:
            pdf_path, cached = self.download_pdf(event)
            archive.write(pdf_path, f"PDF/{prefix}.pdf")
            text_written = False
            if include_text:
                parsed = self.extract_text_cached(pdf_path)
                archive.writestr(f"TXT/{prefix}.txt", parsed["text"])
                text_written = True
            return {"id": event.get("id"), "announcement_id": announcement_id, "title": event.get("title"),
                    "status": "success", "cached": cached, "text": text_written}
        except Exception as exc:  # One malformed PDF must not discard the remaining archive.
            return {"id": event.get("id"), "announcement_id": announcement_id, "title": event.get("title"),
                    "status": "failed", "error": f"{type(exc).__name__}: {str(exc)[:300]}"}

    def download_pdf(self, event: Dict[str, Any]) -> Tuple[Path, bool]:
        url = str(event.get("url") or "")
        parsed = urlparse(url)
        if parsed.scheme != "https" or parsed.hostname != "static.cninfo.com.cn" or not parsed.path.lower().endswith(".pdf"):
            raise AnnouncementArtifactError("公告 PDF 地址不在允许的巨潮 HTTPS 域名")
        metrics = event.get("metrics") or {}
        announcement_id = self._safe_name(str(event.get("external_id") or metrics.get("announcement_id") or "announcement"), 80)
        symbol = self._safe_name(str((event.get("symbols") or [metrics.get("code") or "unknown"])[0]), 20)
        directory = self.root / symbol
        directory.mkdir(parents=True, exist_ok=True)
        target = directory / f"{announcement_id}.pdf"
        if target.exists() and target.stat().st_size > 4:
            with target.open("rb") as existing:
                if existing.read(4) == b"%PDF":
                    return target, True
        # A fixed ``.part`` name is unsafe because the research queue can run
        # two projects for the same company at once.  Keep each writer fully
        # isolated, then publish atomically; concurrent writers download the
        # same immutable CNInfo object, so the last atomic replace is benign.
        handle = tempfile.NamedTemporaryFile(
            prefix=f".{target.name}.",
            suffix=".part",
            dir=directory,
            delete=False,
        )
        temporary = Path(handle.name)
        handle.close()
        try:
            with self.session.get(url, headers={"User-Agent": "Mozilla/5.0", "Referer": "https://www.cninfo.com.cn/"},
                                  timeout=(10, 60), stream=True) as response:
                response.raise_for_status()
                declared = int(response.headers.get("Content-Length") or 0)
                if declared > self.max_pdf_bytes:
                    raise AnnouncementArtifactError("公告 PDF 超过单文件大小上限")
                size = 0
                with temporary.open("wb") as output:
                    for chunk in response.iter_content(64 * 1024):
                        if not chunk:
                            continue
                        size += len(chunk)
                        if size > self.max_pdf_bytes:
                            raise AnnouncementArtifactError("公告 PDF 超过单文件大小上限")
                        output.write(chunk)
                    output.flush()
                    os.fsync(output.fileno())
            with temporary.open("rb") as downloaded:
                if downloaded.read(4) != b"%PDF":
                    raise AnnouncementArtifactError("下载内容不是有效 PDF")
            os.replace(temporary, target)
            return target, False
        finally:
            temporary.unlink(missing_ok=True)

    @classmethod
    def extract_text(cls, pdf_path: Path) -> str:
        return str(cls._extract_text_audited(pdf_path)["text"])

    @classmethod
    def _extract_text_with_metadata(cls, pdf_path: Path) -> Tuple[str, int]:
        """Compatibility wrapper retained for existing callers and adapters."""

        parsed = cls._extract_text_audited(pdf_path)
        return str(parsed["text"]), int(parsed["page_count"])

    @classmethod
    def _extract_text_audited(cls, pdf_path: Path) -> Dict[str, Any]:
        """Run fast extraction out of process with bounded, explicit fallback."""

        try:
            return cls._run_text_extractor(pdf_path, engine="pdfium", timeout=cls._PDFIUM_TIMEOUT_SECONDS)
        except _PdfExtractionFailure as primary_exc:
            if primary_exc.code not in cls._FALLBACK_SAFE_FAILURES:
                raise
            try:
                fallback = cls._run_text_extractor(
                    pdf_path,
                    engine="pypdf",
                    timeout=cls._PYPDF_FALLBACK_TIMEOUT_SECONDS,
                )
                fallback["fallback_reason"] = primary_exc.code
                return fallback
            except _PdfExtractionFailure as fallback_exc:
                raise _PdfExtractionFailure(
                    "fallback_failed",
                    "PDF 文本提取失败"
                    f"（PDFium：{primary_exc.code}；pypdf 降级：{fallback_exc.code}）"
                ) from fallback_exc

    @classmethod
    def _run_text_extractor(cls, pdf_path: Path, *, engine: str, timeout: int) -> Dict[str, Any]:
        path = Path(pdf_path).resolve()
        result_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=f".{path.name}.{engine}.",
                suffix=".extract.json",
                delete=False,
            ) as output:
                result_path = Path(output.name)
            worker_env = os.environ.copy()
            worker_env["DSA_PDF_WORKER_MAX_TEXT_CHARS"] = str(cls._MAX_EXTRACTED_TEXT_CHARS)
            try:
                completed = subprocess.run(
                    [
                        sys.executable,
                        str(Path(__file__).resolve()),
                        "--pdf-text-worker",
                        engine,
                        str(path),
                        str(result_path),
                    ],
                    check=False,
                    env=worker_env,
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                    timeout=timeout,
                )
            except subprocess.TimeoutExpired as exc:
                raise _PdfExtractionFailure("timeout", f"{engine} PDF 文本提取超时（{timeout} 秒）") from exc
            try:
                payload = json.loads(result_path.read_text(encoding="utf-8"))
            except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
                raise _PdfExtractionFailure(
                    "worker_crashed",
                    f"{engine} PDF 提取进程未返回有效结果",
                ) from exc
            if completed.returncode != 0 or payload.get("status") != "success":
                error_code = str(payload.get("error_code") or "worker_crashed")
                raise _PdfExtractionFailure(error_code, f"{engine} PDF 文本提取失败：{error_code}")
            cls._validate_extraction_payload(payload, engine=engine)
            payload.pop("status", None)
            return payload
        finally:
            if result_path is not None:
                result_path.unlink(missing_ok=True)

    @classmethod
    def _validate_extraction_payload(cls, payload: Dict[str, Any], *, engine: str) -> None:
        text = payload.get("text")
        page_count = payload.get("page_count")
        pages_extracted = payload.get("pages_extracted")
        expected_method = "pdfium_text" if engine == "pdfium" else "pypdf_text_fallback"
        expected_status = "complete" if text else "complete_no_selectable_text"
        if (
            not isinstance(text, str)
            or len(text) > cls._MAX_EXTRACTED_TEXT_CHARS
            or not isinstance(page_count, int)
            or page_count <= 0
            or not isinstance(pages_extracted, int)
            or pages_extracted != page_count
            or payload.get("extraction_complete") is not True
            or payload.get("extraction_status") != expected_status
            or payload.get("extraction_method") != expected_method
            or not isinstance(payload.get("extraction_engine_version"), str)
            or not payload.get("extraction_engine_version")
            or not isinstance(payload.get("duration_ms"), int)
            or payload.get("duration_ms") < 0
        ):
            raise _PdfExtractionFailure("incomplete", f"{engine} PDF 提取结果完整性校验失败")

    def fetch_text(self, event: Dict[str, Any]) -> Dict[str, Any]:
        """Download one CNInfo PDF and return its hash-bound cached text."""

        pdf_path, pdf_cached = self.download_pdf(event)
        parsed = self.extract_text_cached(pdf_path)
        return {
            **parsed,
            "pdf_path": pdf_path,
            "pdf_cached": pdf_cached,
        }

    def extract_text_cached(self, pdf_path: Path) -> Dict[str, Any]:
        """Return extracted text cached by the immutable PDF content hash.

        Announcement ids normally identify immutable CNInfo PDFs, but the cache
        key deliberately includes the file hash rather than trusting the id.  A
        replaced or repaired local PDF therefore cannot reuse stale text.
        """

        path = Path(pdf_path).resolve()
        if not self._is_pdf(path):
            raise AnnouncementArtifactError("公告缓存文件不是有效 PDF")
        document_hash = self._file_hash(path)
        document_bytes = path.stat().st_size
        cache_path = path.with_name(f"{path.name}.{document_hash}.text.json")
        with self._text_cache_lock:
            cached = self._read_text_cache(
                cache_path,
                document_hash=document_hash,
                document_bytes=document_bytes,
            )
            if cached is not None:
                return {**cached, "cached": True, "cache_path": cache_path}

            extracted = self._extract_text_audited(path)
            text = str(extracted["text"])
            payload = {
                "version": self._TEXT_CACHE_VERSION,
                "document_hash": document_hash,
                "document_bytes": document_bytes,
                "text": text,
                "text_hash": sha256(text.encode("utf-8", errors="ignore")).hexdigest(),
                "text_chars": len(text),
                "page_count": int(extracted["page_count"]),
                "pages_extracted": int(extracted["pages_extracted"]),
                "extraction_complete": bool(extracted["extraction_complete"]),
                "extraction_status": str(extracted["extraction_status"]),
                "extraction_method": str(extracted["extraction_method"]),
                "extraction_engine_version": str(extracted["extraction_engine_version"]),
                "extraction_duration_ms": int(extracted.get("duration_ms") or 0),
                "fallback_reason": str(extracted.get("fallback_reason") or ""),
            }
            self._write_text_cache(cache_path, payload)
            return {**payload, "cached": False, "cache_path": cache_path}

    def _read_text_cache(
        self,
        cache_path: Path,
        *,
        document_hash: str,
        document_bytes: int,
    ) -> Dict[str, Any] | None:
        if not cache_path.is_file():
            return None
        try:
            payload = json.loads(cache_path.read_text(encoding="utf-8"))
            text = payload.get("text")
            cached_text_chars = payload.get("text_chars")
            cache_version = payload.get("version")
            extraction_method = payload.get("extraction_method")
            if (
                cache_version not in {self._LEGACY_TEXT_CACHE_VERSION, self._TEXT_CACHE_VERSION}
                or payload.get("document_hash") != document_hash
                or int(payload.get("document_bytes") or 0) != document_bytes
                or not isinstance(text, str)
                or not isinstance(cached_text_chars, int)
                or cached_text_chars != len(text)
                or payload.get("text_hash") != sha256(text.encode("utf-8", errors="ignore")).hexdigest()
                or extraction_method not in self._SUPPORTED_EXTRACTION_METHODS
            ):
                return None
            if cache_version == self._TEXT_CACHE_VERSION:
                page_count = payload.get("page_count")
                expected_status = "complete" if text else "complete_no_selectable_text"
                if (
                    not isinstance(page_count, int)
                    or page_count <= 0
                    or payload.get("pages_extracted") != page_count
                    or payload.get("extraction_complete") is not True
                    or payload.get("extraction_status") != expected_status
                    or not isinstance(payload.get("extraction_engine_version"), str)
                    or not payload.get("extraction_engine_version")
                    or not isinstance(payload.get("extraction_duration_ms"), int)
                    or payload.get("extraction_duration_ms") < 0
                ):
                    return None
            else:
                # v1 was written only after pypdf iterated every page without an
                # exception, so it can be safely normalized without an expensive
                # one-time reparse.
                page_count = int(payload.get("page_count") or 0)
                if page_count <= 0:
                    return None
                payload = {
                    **payload,
                    "pages_extracted": page_count,
                    "extraction_complete": True,
                    "extraction_status": "complete" if text else "complete_no_selectable_text",
                    "extraction_engine_version": "pypdf=legacy-cache-v1",
                    "extraction_duration_ms": 0,
                    "fallback_reason": "legacy_cache",
                }
            return payload
        except (OSError, TypeError, ValueError, json.JSONDecodeError):
            return None

    @staticmethod
    def _write_text_cache(cache_path: Path, payload: Dict[str, Any]) -> None:
        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                mode="w",
                encoding="utf-8",
                prefix=f"{cache_path.name}.",
                suffix=".part",
                dir=cache_path.parent,
                delete=False,
            ) as output:
                temporary_path = Path(output.name)
                json.dump(payload, output, ensure_ascii=False, separators=(",", ":"))
                output.flush()
                os.fsync(output.fileno())
            temporary_path.replace(cache_path)
            temporary_path = None
        except OSError as exc:
            raise AnnouncementArtifactError(f"公告正文缓存写入失败：{type(exc).__name__}") from exc
        finally:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)

    @staticmethod
    def _is_pdf(path: Path) -> bool:
        if not path.is_file() or path.stat().st_size <= 4:
            return False
        try:
            with path.open("rb") as handle:
                return handle.read(4) == b"%PDF"
        except OSError:
            return False

    @staticmethod
    def _file_hash(path: Path) -> str:
        digest = sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(128 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _safe_name(value: str, limit: int) -> str:
        return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff._-]+", "_", value).strip("._")[:limit] or "announcement"


if __name__ == "__main__":
    raise SystemExit(_run_pdf_text_worker_cli(sys.argv))
